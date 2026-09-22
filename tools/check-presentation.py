#!/usr/bin/env python3
"""Check PDF slide images with independent PDF and PresentationML readers."""
import argparse
from datetime import datetime, timezone
from io import BytesIO
import hashlib
import json
from pathlib import Path
import platform
import posixpath
import random
import re
import statistics
import subprocess
import xml.etree.ElementTree as ET
import zipfile

from PIL import Image, ImageChops, ImageStat
from pptx import Presentation
from pypdf import PdfReader, PdfWriter
from pypdf.generic import NameObject, NumberObject, RectangleObject
import pypdfium2 as pdfium
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader


def main():
    root = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser()
    parser.add_argument("--command", type=Path, default=root / ".build/debug/allomer")
    parser.add_argument("--benchmark", action="store_true")
    args = parser.parse_args()
    args.command = args.command.resolve()
    work = root / ".tools/presentation-check"
    work.mkdir(exist_ok=True)
    env = {"PATH": "/usr/bin:/bin"}
    if ".app/" not in str(args.command):
        env["ALLOMER_TOOLS_DIR"] = str(root / ".tools/bin")

    raw = work / "authored.pdf"
    pdf = canvas.Canvas(str(raw), pageCompression=1, invariant=1)
    for number, (width, height) in enumerate([(360, 240), (240, 360), (400, 300), (20, 40)], 1):
        pdf.setPageSize((width, height))
        pdf.setFillColorRGB(0.05, 0.25, 0.5)
        pdf.rect(0, 0, width / 3, height / 2, fill=1, stroke=0)
        pdf.setFillColorRGB(0.8, 0.2, 0.1)
        pdf.circle(width * 0.7, height * 0.6, min(width, height) / 8, fill=1, stroke=0)
        pdf.setFillAlpha(0.4)
        pdf.setFillColorRGB(0.2, 0.9, 0.3)
        pdf.rect(width / 2, height / 2, width / 4, height / 4, fill=1, stroke=0)
        pdf.setFillAlpha(1)
        pdf.setFillColorRGB(0, 0, 0)
        pdf.setFont("Helvetica", min(18, width / 12))
        pdf.drawString(width / 10, height * 0.8, f"Original page {number}")
        pdf.showPage()
    pdf.save()
    writer = PdfWriter(clone_from=raw)
    special = writer.pages[2]
    special.cropbox = RectangleObject([20, 30, 320, 230])
    special.rotate(90)
    special[NameObject("/UserUnit")] = NumberObject(2)
    source = work / "original café 100%.pdf"
    writer.write(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    output = work / "pages.pptx"
    output.unlink(missing_ok=True)
    result = subprocess.run([args.command, "convert", source, output], cwd=work, env=env,
                            capture_output=True, text=True, timeout=130)
    assert result.returncode == 0, result.stderr
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before

    deck = Presentation(output)
    assert len(deck.slides) == 4
    assert (deck.slide_width, deck.slide_height) == (360 * 12700, 240 * 12700)
    document = pdfium.PdfDocument(source)
    source_pages = PdfReader(source).pages
    for index, slide in enumerate(deck.slides):
        assert len(slide.shapes) == 1
        shape = slide.shapes[0]
        image = Image.open(BytesIO(shape.image.blob))
        assert image.format == "PNG" and image.mode == "RGB"
        # PDFium exposes box coordinates in PDF units; UserUnit also scales the physical page.
        unit = float(source_pages[index].get("/UserUnit", 1))
        bitmap = document[index].render(scale=2 * unit)
        expected = bitmap.to_pil().convert("RGB")
        assert image.size == expected.size, (index, image.size, expected.size)
        difference = ImageChops.difference(image, expected)
        assert max(ImageStat.Stat(difference).mean) < 3, (index, ImageStat.Stat(difference).mean)
        assert max(ImageStat.Stat(image).stddev) > 10, "A slide image is blank."
        assert abs(shape.width / shape.height - image.width / image.height) < 0.001
        assert abs(shape.left * 2 + shape.width - deck.slide_width) <= 1
        assert abs(shape.top * 2 + shape.height - deck.slide_height) <= 1
        assert shape.left >= 0 and shape.top >= 0
        assert shape.left + shape.width <= deck.slide_width and shape.top + shape.height <= deck.slide_height
        image.save(work / f"page-{index + 1}.png")
        expected.save(work / f"expected-{index + 1}.png")
        bitmap.close()
    document.close()

    with zipfile.ZipFile(output) as archive:
        assert archive.testzip() is None
        names = set(archive.namelist())
        assert len(names) == len(archive.namelist())
        for name in names:
            if name.endswith((".xml", ".rels")):
                tree = ET.fromstring(archive.read(name))
                if name.endswith(".rels"):
                    folder = posixpath.dirname(posixpath.dirname(name))
                    ids = set()
                    for relation in tree:
                        assert relation.attrib["Id"] not in ids
                        ids.add(relation.attrib["Id"])
                        assert relation.attrib.get("TargetMode", "Internal") == "Internal"
                        target = posixpath.normpath(posixpath.join(folder, relation.attrib["Target"]))
                        assert target in names, (name, target)
            if name.endswith(".png"):
                assert archive.getinfo(name).compress_type == zipfile.ZIP_STORED

    original_output = output.read_bytes()
    result = subprocess.run([args.command, "convert", source, output], cwd=work, env=env, capture_output=True, timeout=130)
    assert result.returncode != 0 and output.read_bytes() == original_output
    bad = work / "invalid.pdf"
    bad.write_bytes(b"%PDF-1.7\nNot a document")
    refused = work / "invalid.pptx"
    refused.unlink(missing_ok=True)
    result = subprocess.run([args.command, "convert", bad, refused], cwd=work, env=env, capture_output=True, timeout=130)
    assert result.returncode != 0 and not refused.exists()
    for label, page, unit, succeeds in [("small-canvas", source_pages[3], 1, True),
                                        ("large-canvas", source_pages[0], 12, True),
                                        ("bad-unit", source_pages[0], 0, False),
                                        ("pixel-limit", source_pages[0], 100, False)]:
        writer = PdfWriter()
        copied = writer.add_page(page)
        copied[NameObject("/UserUnit")] = NumberObject(unit)
        fixture = work / f"{label}.pdf"
        writer.write(fixture)
        converted = work / f"{label}.pptx"
        converted.unlink(missing_ok=True)
        result = subprocess.run([args.command, "convert", fixture, converted], cwd=work, env=env,
                                capture_output=True, text=True, timeout=130)
        assert (result.returncode == 0) == succeeds, (label, result.stderr)
        assert converted.exists() == succeeds
        if succeeds:
            deck = Presentation(converted)
            assert 914400 <= deck.slide_width <= 51206400 and 914400 <= deck.slide_height <= 51206400
            shape = deck.slides[0].shapes[0]
            assert abs(shape.width / shape.height - float(page.cropbox.width / page.cropbox.height)) < 0.001
            assert shape.left >= 0 and shape.top >= 0
    assert not any(path.name.startswith((".allomer-", "zip-part-")) for path in work.iterdir())
    print(f"Verified four slides, crop, rotation, UserUnit, transparency, picture placement, ZIP relationships, canvas and pixel limits, and no overwrite: {output}")
    if args.benchmark:
        assert ".app/" in str(args.command), "Use a packaged release build for the benchmark."
        texture = Image.frombytes("RGB", (1024, 1024), random.Random(42).randbytes(1024 * 1024 * 3))
        groups = []
        for count in (1, 12):
            source = work / f"benchmark-{count}.pdf"
            pdf = canvas.Canvas(str(source), pagesize=(595.28, 841.89), invariant=1, pageCompression=1)
            for index in range(count):
                pdf.setFont("Helvetica-Bold", 24)
                pdf.drawString(36, 790, f"Original slide benchmark {index + 1}")
                pdf.drawImage(ImageReader(texture), 36, 100, width=520, height=520)
                pdf.setFillColorRGB(0.1, 0.3, 0.6)
                pdf.rect(36, 650, 520, 80, fill=1, stroke=0)
                pdf.showPage()
            pdf.save()
            samples = []
            for _ in range(3):
                output = work / f"benchmark-{count}.pptx"
                output.unlink(missing_ok=True)
                result = subprocess.run(["/usr/bin/time", "-l", args.command, "convert", source, output],
                    cwd=work, env=env, check=True, capture_output=True, text=True, timeout=300)
                elapsed = re.search(r"([0-9.]+) real", result.stderr)
                memory = re.search(r"([0-9]+)\s+maximum resident set size", result.stderr)
                assert elapsed and memory and len(Presentation(output).slides) == count, result.stderr
                samples.append({"elapsed_seconds": float(elapsed[1]), "reported_peak_resident_bytes": int(memory[1]),
                                "output_bytes": output.stat().st_size})
            groups.append({"pages": count, "source_bytes": source.stat().st_size,
                "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "samples": samples,
                "median_elapsed_seconds": statistics.median(s["elapsed_seconds"] for s in samples),
                "median_reported_peak_resident_bytes": statistics.median(s["reported_peak_resident_bytes"] for s in samples)})
        helpers = args.command.parent.parent / "Helpers"
        report = {"recorded_at_utc": datetime.now(timezone.utc).isoformat(), "macos": platform.mac_ver()[0],
            "architecture": platform.machine(), "resolution_dpi": 144,
            "workload": "One and twelve A4 pages, each with text, a vector shape, and a deterministic 1024-square RGB noise image",
            "command_sha256": hashlib.sha256(args.command.read_bytes()).hexdigest(),
            "helper_sha256": hashlib.sha256((helpers / "mutool").read_bytes()).hexdigest(),
            "launcher_sha256": hashlib.sha256((helpers / "pdfguard").read_bytes()).hexdigest(),
            "memory_scope": "Maximum resident size reported by macOS time -l for the conversion command. This is not the sum of simultaneously resident parent and helper processes. The GUI is excluded.",
            "timing_scope": "Complete command conversion, including page rendering, ZIP writing, content checks, and publication. Warm local runs; not a comparison with another app.",
            "runs": groups}
        (root / "research/presentation-performance.json").write_text(json.dumps(report, indent=2) + "\n")
        print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
