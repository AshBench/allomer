#!/usr/bin/env python3
"""Check original PowerPoint fixtures with an independent PDF reader and rasterizer."""
import argparse
from datetime import datetime, timezone
from io import BytesIO
import hashlib
import json
from pathlib import Path
import platform
import random
import re
import statistics
import subprocess
import unicodedata
from urllib.parse import quote
import xml.etree.ElementTree as ET
import zipfile

from PIL import Image, ImageStat
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from pypdf import PdfReader
import pypdfium2 as pdfium


def fixture(path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(7.5)
    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = "Original inherited title"
    first.placeholders[1].text = "First body paragraph\nSecond body paragraph"
    for shape in first.shapes:
        for paragraph in shape.text_frame.paragraphs:
            paragraph.font.name = "Arial"
    box = first.shapes.add_textbox(Inches(1), Inches(5.5), Inches(8), Inches(1))
    for text, color, bold, size in [("Red span ", (210, 20, 20), True, 26), ("Blue span", (20, 20, 210), False, 18)]:
        run = box.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.name, run.font.size, run.font.bold = "Arial", Pt(size), bold
        run.font.color.rgb = RGBColor(*color)
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.background.fill.solid()
    second.background.fill.fore_color.rgb = RGBColor(255, 250, 210)
    group = second.shapes.add_group_shape()
    for index, color in enumerate([(10, 150, 100), (70, 90, 220)]):
        rectangle = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5 + index), Inches(0.5), Inches(0.8), Inches(0.8))
        rectangle.fill.solid()
        rectangle.fill.fore_color.rgb = RGBColor(*color)
        rectangle.line.fill.background()
    group.rotation = 25
    picture = Image.new("RGB", (128, 64))
    picture.putdata([(x * 2, y * 4, (x + y) % 256) for y in range(64) for x in range(128)])
    buffer = BytesIO()
    picture.save(buffer, format="PNG")
    image = second.shapes.add_picture(buffer, Inches(3), Inches(0.5), width=Inches(3), height=Inches(1.5))
    image.crop_left, image.crop_top = 0.15, 0.1
    table = second.shapes.add_table(2, 2, Inches(0.5), Inches(2.2), Inches(5), Inches(1.3)).table
    for row in range(2):
        for column in range(2):
            cell = table.cell(row, column)
            cell.text = f"Cell {row + 1}-{column + 1}"
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(20, 50, 100) if row == 0 else RGBColor(235, 215, 215)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name, paragraph.font.size = "Arial", Pt(18)
                paragraph.font.color.rgb = RGBColor(255, 255, 255) if row == 0 else RGBColor(160, 20, 20)
    data = CategoryChartData()
    data.categories = ["North", "South", "West"]
    data.add_series("Original counts", (2, 4, 1))
    chart = second.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(3.7), Inches(8.5), Inches(3.2), data).chart
    chart.has_legend, chart.chart_style = False, 10
    deck.save(path)


def main():
    root = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser()
    parser.add_argument("--command", type=Path, default=root / ".build/debug/allomer")
    parser.add_argument("--benchmark", action="store_true")
    args = parser.parse_args()
    command = args.command.resolve()
    work = root / ".tools/presentation-reader-check"
    work.mkdir(exist_ok=True)
    env = {"PATH": "/usr/bin:/bin"}
    if ".app/" not in str(command):
        env["ALLOMER_TOOLS_DIR"] = str(root / ".tools/bin")

    def convert(source, output, succeeds=True):
        output.unlink(missing_ok=True)
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        result = subprocess.run([command, "convert", source, output], cwd=work, env=env,
                                capture_output=True, text=True, timeout=130)
        assert (result.returncode == 0) == succeeds, (source.name, result.stderr)
        assert output.exists() == succeeds, source.name
        assert hashlib.sha256(source.read_bytes()).hexdigest() == before
        assert not any(p.name.startswith(".allomer-") for p in work.iterdir())
        return output

    source = work / "original café 100%.pptx"
    fixture(source)
    output = convert(source, work / "original.pdf")
    pdf = PdfReader(output)
    assert len(pdf.pages) == 2
    for page in pdf.pages:
        assert tuple(float(v) for v in page.mediabox) == (0, 0, 720, 540)
    for index, texts in enumerate([["Original inherited title", "First body paragraph", "Second body paragraph", "Red span", "Blue span"],
                                  ["Cell 1-1", "Cell 1-2", "Cell 2-1", "Cell 2-2", "North", "South", "West", "Original counts"]]):
        actual = "".join(pdf.pages[index].extract_text().split())
        for text in texts:
            assert "".join(text.split()) in actual, (index, text, actual)
    # WebKit rasterizes the two rotated group shapes. The chart stays in PDF content.
    assert len(pdf.pages[0].images) == 0
    assert sorted(image.image.size for image in pdf.pages[1].images) == [(128, 64), (178, 183), (178, 183)]
    document = pdfium.PdfDocument(output)
    for index in range(2):
        bitmap = document[index].render(scale=1)
        image = bitmap.to_pil().convert("RGB")
        assert max(ImageStat.Stat(image).stddev) > 20
        if index == 1:
            assert all(abs(a - b) <= 2 for a, b in zip(image.getpixel((700, 20)), (255, 250, 210)))
            assert all(abs(a - b) <= 2 for a, b in zip(image.getpixel((40, 162)), (20, 50, 100)))
        image.save(work / f"slide-{index + 1}.png")
        bitmap.close()
    document.close()
    original = output.read_bytes()
    result = subprocess.run([command, "convert", source, output], env=env, capture_output=True, timeout=130)
    assert result.returncode != 0 and output.read_bytes() == original

    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}

    def changed(label, updates, succeeds=False):
        file = work / f"{label}.pptx"
        with zipfile.ZipFile(file, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in (parts | updates).items():
                if data is not None:
                    archive.writestr(name, data)
        return convert(file, work / f"{label}.pdf", succeeds)

    pml = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    presentation = ET.fromstring(parts["ppt/presentation.xml"])
    slides = presentation.find(pml + "sldIdLst")
    slides[:] = list(reversed(list(slides)))
    hidden = ET.fromstring(parts["ppt/slides/slide2.xml"])
    hidden.set("show", "0")
    reordered = changed("reordered-hidden", {"ppt/presentation.xml": ET.tostring(presentation),
                                             "ppt/slides/slide2.xml": ET.tostring(hidden)}, True)
    result = PdfReader(reordered)
    assert len(result.pages) == 2
    assert "Cell1-1" in "".join(result.pages[0].extract_text().split())
    assert "inherited" in result.pages[1].extract_text()
    changed("missing-slide", {"ppt/slides/slide2.xml": None})
    media = next(name for name in parts if name.startswith("ppt/media/"))
    relpath = "ppt/slides/_rels/slide2.xml.rels"
    encoded = ET.fromstring(parts[relpath])
    for relationship in encoded:
        if relationship.get("Type", "").endswith("/image"):
            relationship.set("Target", "../media/caf%C3%A9%20100%25.png")
    changed("encoded-media", {media: None, "ppt/media/café 100%.png": parts[media], relpath: ET.tostring(encoded)}, True)
    decomposed = unicodedata.normalize("NFD", "café 100%.png")
    for relationship in encoded:
        if relationship.get("Type", "").endswith("/image"):
            relationship.set("Target", "../media/" + quote(decomposed))
    changed("decomposed-media", {media: None, "ppt/media/" + decomposed: parts[media], relpath: ET.tostring(encoded)}, True)
    for form in ("NFC", "NFD"):
        name = unicodedata.normalize(form, "café 100%.xml")
        relations = ET.fromstring(parts["ppt/_rels/presentation.xml.rels"])
        for relationship in relations:
            if relationship.get("Target") == "slides/slide2.xml":
                relationship.set("Target", "slides/" + quote(name))
        types = ET.fromstring(parts["[Content_Types].xml"])
        for content_type in types:
            if content_type.get("PartName") == "/ppt/slides/slide2.xml":
                content_type.set("PartName", "/ppt/slides/" + quote(name))
        renamed = changed("unicode-slide-" + form, {"ppt/slides/slide2.xml": None, relpath: None,
            "ppt/slides/" + name: parts["ppt/slides/slide2.xml"], "ppt/slides/_rels/" + name + ".rels": parts[relpath],
            "ppt/_rels/presentation.xml.rels": ET.tostring(relations), "[Content_Types].xml": ET.tostring(types)}, True)
        assert "Cell1-1" in "".join(PdfReader(renamed).pages[1].extract_text().split())
    changed("empty-unused-part", {"unused.bin": b""}, True)
    changed("missing-image", {media: None})
    changed("broken-image", {media: b"not an image"})
    changed("broken-xml", {"ppt/slides/slide2.xml": b"<broken"})
    changed("entity", {"ppt/slides/slide2.xml": b'<!DOCTYPE s [<!ENTITY x "expanded">]><s>&x;</s>'})
    changed("path-traversal", {"../outside": b"blocked"})
    changed("part-limit", {"large.bin": bytes(32 * 1024 * 1024 + 1)})
    changed("wrong-root", {"ppt/presentation.xml": b'<presentation><sldIdLst><sldId/></sldIdLst></presentation>'})
    large = ET.fromstring(parts["ppt/presentation.xml"])
    large.find(pml + "sldSz").set("cx", "999999999")
    changed("canvas-limit", {"ppt/presentation.xml": ET.tostring(large)})
    relationships = ET.fromstring(parts[relpath])
    for relationship in relationships:
        if relationship.get("Type", "").endswith("/image"):
            relationship.set("TargetMode", "External")
            relationship.set("Target", "https://example.invalid/never-download.png")
    changed("network-image", {relpath: ET.tostring(relationships)})
    tools = command.parent.parent / "Helpers" if ".app/" in str(command) else root / ".tools/bin"
    resources = (tools / "webguard").resolve().parent.parent / "Resources/Presentation"
    manifest = work / "media.json"
    refused = work / "invalid-media-index.pdf"
    try:
        for label in ("../original.pdf", "/etc/passwd", "0", "01", "100001"):
            manifest.write_text(json.dumps({"ppt/media/image.png": label}))
            refused.unlink(missing_ok=True)
            result = subprocess.run([tools / "webguard", source, work, resources, "pptx", source, refused.name, "2"],
                                    cwd=work, env=env, capture_output=True, text=True, timeout=130)
            assert result.returncode != 0 and "media index is invalid" in result.stderr, result.stderr
            assert not refused.exists()
        link = work / "99999"
        link.unlink(missing_ok=True)
        link.symlink_to(source)
        try:
            manifest.write_text(json.dumps({"ppt/media/image.png": link.name}))
            result = subprocess.run([tools / "webguard", source, work, resources, "pptx", source, refused.name, "2"],
                                    cwd=work, env=env, capture_output=True, text=True, timeout=130)
            assert result.returncode != 0 and "regular file" in result.stderr, result.stderr
            assert not refused.exists()
        finally:
            link.unlink()
    finally:
        manifest.unlink(missing_ok=True)
    print(f"Verified slide order, hidden slides, text, vector charts, raster colors, damaged parts, limits, offline resources, cleanup, and no overwrite: {output}")
    if args.benchmark:
        assert ".app/" in str(command), "Use a packaged release build for the benchmark."
        texture = Image.frombytes("RGB", (1024, 1024), random.Random(42).randbytes(1024 * 1024 * 3))
        picture = BytesIO()
        texture.save(picture, format="PNG")
        groups = []
        for count, distinct in ((1, False), (12, False), (12, True)):
            deck = Presentation()
            deck.slide_width, deck.slide_height = Inches(10), Inches(7.5)
            for index in range(count):
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                text = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1)).text_frame
                text.text = f"Original reader benchmark {index + 1}"
                text.paragraphs[0].font.name, text.paragraphs[0].font.size = "Arial", Pt(24)
                if distinct:
                    pixels = random.Random(42 + index).randbytes(1024 * 1024 * 3)
                    picture = BytesIO()
                    Image.frombytes("RGB", (1024, 1024), pixels).save(picture, format="PNG")
                slide.shapes.add_picture(BytesIO(picture.getvalue()), Inches(2), Inches(1.5), Inches(5.5), Inches(5.5))
            label = f"benchmark-{count}" + ("-distinct" if distinct else "")
            source = work / f"{label}.pptx"
            deck.save(source)
            samples = []
            for _ in range(3):
                output = work / f"{label}.pdf"
                output.unlink(missing_ok=True)
                result = subprocess.run(["/usr/bin/time", "-l", command, "convert", source, output], cwd=work, env=env,
                                        check=True, capture_output=True, text=True, timeout=130)
                elapsed = re.search(r"([0-9.]+) real", result.stderr)
                memory = re.search(r"([0-9]+)\s+maximum resident set size", result.stderr)
                assert elapsed and memory, result.stderr
                pdf = PdfReader(output)
                assert len(pdf.pages) == count
                for index, page in enumerate(pdf.pages):
                    assert f"Original reader benchmark {index + 1}" in page.extract_text()
                samples.append({"elapsed_seconds": float(elapsed[1]), "reported_peak_resident_bytes": int(memory[1]),
                                "output_bytes": output.stat().st_size})
            groups.append({"slides": count, "distinct_images": distinct, "source_bytes": source.stat().st_size,
                "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "samples": samples,
                "median_elapsed_seconds": statistics.median(s["elapsed_seconds"] for s in samples),
                "median_reported_peak_resident_bytes": statistics.median(s["reported_peak_resident_bytes"] for s in samples)})
        contents = command.parent.parent
        report = {"recorded_at_utc": datetime.now(timezone.utc).isoformat(), "macos": platform.mac_ver()[0],
            "architecture": platform.machine(), "workload": "One and twelve 10 by 7.5 inch slides, each with text and a deterministic 1024-square RGB noise image. Twelve-slide cases cover both one shared image and twelve distinct images.",
            "sha256": {str(path.relative_to(contents)): hashlib.sha256(path.read_bytes()).hexdigest() for path in
                [command, contents / "Helpers/webconvert", contents / "Helpers/webguard", contents / "Resources/Presentation/renderer.js"]},
            "memory_scope": "macOS time -l reports the conversion command and its waited-for child processes. This is not aggregate memory. It excludes separate WebKit services and the GUI.",
            "timing_scope": "Complete command conversion, including ZIP and image validation, WebKit rendering, PDF cleanup, checks, and publication. Three warm local runs; no cross-app comparison.",
            "runs": groups}
        (root / "research/presentation-reader-performance.json").write_text(json.dumps(report, indent=2) + "\n")
        print(json.dumps(groups, indent=2))


if __name__ == "__main__":
    main()
